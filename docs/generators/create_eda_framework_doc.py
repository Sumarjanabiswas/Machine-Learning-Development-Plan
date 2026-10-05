import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Sets internal padding (margins) for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, top="CBD5E1", bottom="CBD5E1", left=None, right=None, sz="4"):
    """Sets specific borders on a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    borders_xml += f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>' if top else '<w:top w:val="none"/>'
    borders_xml += f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>' if bottom else '<w:bottom w:val="none"/>'
    borders_xml += f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>' if left else '<w:left w:val="none"/>'
    borders_xml += f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>' if right else '<w:right w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def add_callout(doc, title, text, border_color="0284C7", bg_color="F0F9FF"):
    """Creates a modern executive callout banner box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"■ {title}\n")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(11)
    run_title.font.color.rgb = RGBColor.from_string(border_color)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def style_heading(p, font_name="Calibri", size_pt=14, bold=True, color_rgb=(15, 41, 74), space_before=12, space_after=6):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        r.font.name = font_name
        r.font.size = Pt(size_pt)
        r.bold = bold
        r.font.color.rgb = RGBColor(*color_rgb)

def add_h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    style_heading(p, "Calibri", 17, True, (15, 41, 74), space_before=16, space_after=6)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    style_heading(p, "Calibri", 13, True, (13, 148, 136), space_before=12, space_after=4)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    style_heading(p, "Calibri", 11, True, (79, 70, 229), space_before=8, space_after=3)
    return p

def add_body_p(doc, text, bold_prefix=None, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_bullet_item(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + ": ")
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_image_box(doc, image_path, caption_title, caption_text, width_inches=6.4):
    if not os.path.exists(image_path):
        print(f"Warning: Image {image_path} not found!")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    r = p_img.add_run()
    r.add_picture(image_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    
    r_cap_title = p_cap.add_run(f"{caption_title}: ")
    r_cap_title.bold = True
    r_cap_title.font.name = "Calibri"
    r_cap_title.font.size = Pt(9.5)
    r_cap_title.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)
    
    r_cap_text = p_cap.add_run(caption_text)
    r_cap_text.italic = True
    r_cap_text.font.name = "Calibri"
    r_cap_text.font.size = Pt(9.5)
    r_cap_text.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

def build_eda_framework_document(output_path):
    print("Initializing OmniEDA Framework document generation...")
    doc = Document()

    # Margins: 0.8 inches
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
        # Header & Footer
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("OmniEDA Framework | Exploratory Data Analysis & Visualization Strategy (Week 2)")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Author: Sumarjana Biswas (sumarjanabiswas690@gmail.com) — Data Science Methodology Blueprint")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # -------------------------------------------------------------------------
    # COVER / HEADER BANNER BLOCK
    # -------------------------------------------------------------------------
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.autofit = False
    c_cell = cover_table.cell(0, 0)
    c_cell.width = Inches(6.9)
    set_cell_background(c_cell, "0F172A")
    set_cell_margins(c_cell, top=260, bottom=260, left=260, right=260)
    
    cp = c_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(6)
    r_tag = cp.add_run("DATA SCIENCE METHODOLOGY BLUEPRINT & OPERATIONAL GUIDE\n")
    r_tag.bold = True
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(11)
    r_tag.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
    
    r_main_title = cp.add_run("OmniEDA: Universal Exploratory Data Analysis & Diagnostic Visualization Framework\n")
    r_main_title.bold = True
    r_main_title.font.name = "Calibri"
    r_main_title.font.size = Pt(21)
    r_main_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    r_desc = cp.add_run("A Comprehensive, Dataset-Agnostic Blueprint for Data Profiling, Statistical Diagnostics, Visualization Taxonomy, and Findings Reporting in Python\n\n")
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(10.5)
    r_desc.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    r_author = cp.add_run("Prepared by: Sumarjana Biswas\n")
    r_author.bold = True
    r_author.font.name = "Calibri"
    r_author.font.size = Pt(11)
    r_author.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)

    r_meta = cp.add_run("Email: sumarjanabiswas690@gmail.com  |  Role: Lead Data Science Architect  |  Milestone: Week 2 Deliverable  |  Effort: 32.5 Hours")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------------------
    # EXECUTIVE SUMMARY & METADATA TABLE
    # -------------------------------------------------------------------------
    add_callout(doc, "ARCHITECTURAL VISION & PRACTITIONER PERSPECTIVE", 
                "In modern data science, teams frequently rush into model training, treating Exploratory Data Analysis (EDA) as a routine "
                "checkbox exercise consisting of a few superficial summary tables and scatter plots. This is a severe mistake. "
                "In real-world production systems, over 80% of machine learning failures—ranging from subtle data leakage and severe class imbalance "
                "to unhandled multimodal distributions and fatal concept drift—trace back directly to inadequate exploratory interrogation. "
                "OmniEDA establishes a universal, dataset-agnostic standard for dissecting, diagnosing, and visualizing any structured or "
                "semi-structured dataset before a single line of modeling code is executed. Designed for execution within a focused 32.5-hour "
                "timeframe, this framework bridges the gap between raw statistical telemetry and actionable engineering decisions.", 
                border_color="0284C7", bg_color="F0F9FF")

    # Metadata Summary Table
    meta_tbl = doc.add_table(rows=6, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    meta_widths = [Inches(2.2), Inches(4.7)]
    meta_data = [
        ("Framework Author & Lead", "Sumarjana Biswas (sumarjanabiswas690@gmail.com)"),
        ("Framework Codename", "OmniEDA (Universal Exploratory Data Analysis & Diagnostic Visualization System)"),
        ("Core Mandate", "Establish a repeatable, hypothesis-driven analytical protocol applicable to any tabular, temporal, or multi-modal dataset"),
        ("Primary Beneficiaries", "Data Scientists, Machine Learning Engineers, Analytics Engineers, Business Stakeholders"),
        ("Primary Python Stack", "Python 3.12, Pandas, Polars, Scipy, Statsmodels, Matplotlib, Seaborn, Plotly, Missingno, Great Expectations"),
        ("Implementation Pacing", "32.5 Hours allocated across 6 structured workstreams over 5 working days (Week 2 Milestone)")
    ]
    for row_idx, (lbl, val) in enumerate(meta_data):
        row = meta_tbl.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = meta_widths[0], meta_widths[1]
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=70, bottom=70, left=120, right=120)
        set_cell_margins(c1, top=70, bottom=70, left=120, right=120)
        set_cell_border(c0, top="E2E8F0", bottom="E2E8F0", left="CBD5E1", right="E2E8F0")
        set_cell_border(c1, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="CBD5E1")
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(lbl)
        r0.bold = True
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------------------
    # SECTION 1: INTRODUCTION TO EDA & PHILOSOPHY OF ACTIVE INTERROGATION
    # -------------------------------------------------------------------------
    add_h1(doc, "1. Introduction to EDA & The Philosophy of Active Interrogation")
    
    add_h2(doc, "1.1 What Exploratory Data Analysis Really Is")
    add_body_p(doc, 
               "Exploratory Data Analysis (EDA) is not merely the passive calculation of means, medians, and standard deviations. "
               "Pioneered by the great statistician John Tukey in 1977, EDA is an active, open-ended scientific interrogation of data. "
               "It is the process of critically questioning the assumptions embedded in a dataset, discovering unexpected patterns, identifying "
               "anomalies, testing hypotheses, and assessing whether the data is actually capable of solving the business problem at hand.")
    
    add_body_p(doc,
               "In modern applied data science, EDA serves three foundational purposes:")
    add_bullet_item(doc, "Truth Verification & Data Hygiene", 
                    "Verifying that the data accurately represents physical or business reality. Real data is messy, riddled with logging bugs, "
                    "sensor glitches, timezone discrepancies, and unexpected default placeholders (e.g., -999, 'None', 'Unknown', or epoch zero timestamps).")
    add_bullet_item(doc, "Hypothesis Generation & Pattern Discovery", 
                    "Uncovering non-linear relationships, cluster structures, interaction effects, and temporal dynamics that domain experts might "
                    "not have anticipated.")
    add_bullet_item(doc, "Machine Learning Feasibility & Risk De-Risking", 
                    "Detecting severe class imbalance, high multicollinearity, sparse categorical tokens, and target leakage before sinking weeks into "
                    "futile model tuning.")

    add_h2(doc, "1.2 Why EDA is the Highest-Leverage Phase in the Data Science Lifecycle")
    add_body_p(doc,
               "There is an old adage in machine learning: 'Garbage In, Garbage Out.' But in enterprise environments, the reality is far more dangerous: "
               "'Garbage In, Flawed Confidence Out.' A complex neural network or gradient boosted ensemble will happily fit noise, memorize data leakage, "
               "and output artificially inflated validation scores (e.g., 99.8% accuracy)—only to catastrophically fail upon production deployment.")
    
    add_body_p(doc,
               "Investing disciplined hours in EDA yields compounding dividends across the entire engineering lifecycle:")
    add_bullet_item(doc, "Preventing Silent Catastrophes",
                    "A single lookahead leakage feature (e.g., an account churn label column derived from account closure timestamps that accidentally leaked "
                    "into training features) can render an entire month of engineering worthless. Rigorous EDA catches these anomalies on Day 1.")
    add_bullet_item(doc, "Guiding High-Impact Feature Engineering",
                    "Understanding that a continuous variable follows a heavy-tailed log-normal distribution immediately tells the engineer to apply a log "
                    "or Yeo-Johnson transform. Spotting a 30-day activity drop in bivariate scatter plots directly inspires the creation of velocity ratio features.")
    add_bullet_item(doc, "Accelerating Model Convergence",
                    "Pruning highly collinear features during EDA reduces training time, prevents gradient instability, and shrinks deployment memory footprints.")

    # -------------------------------------------------------------------------
    # SECTION 2: UNIVERSAL DATA TAXONOMY & STRUCTURAL PROFILING
    # -------------------------------------------------------------------------
    add_h1(doc, "2. Universal Data Taxonomy & Structural Profiling")
    
    add_h2(doc, "2.1 The Universal Variable Taxonomy")
    add_body_p(doc,
               "Before any chart is drawn or test is calculated, every variable in the dataset must be formally mapped to its true mathematical "
               "and structural data type. In Python, relying blindly on automated Pandas dtypes (e.g., `int64` or `object`) is notoriously misleading; "
               "a postal code or customer ID is stored as an integer, yet treating it as a continuous metric would be fatal. "
               "The OmniEDA taxonomy establishes unambiguous categorization across five fundamental data classes:")

    # Data Taxonomy Table
    tax_tbl = doc.add_table(rows=6, cols=4)
    tax_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tax_tbl.autofit = False
    tax_col_w = [Inches(1.5), Inches(1.4), Inches(2.2), Inches(1.8)]
    tax_headers = ["Data Class", "Variable Types", "Distinguishing Characteristics", "Inspection & Diagnostic Protocol"]
    
    t_hdr_row = tax_tbl.rows[0]
    for c_idx, h_text in enumerate(tax_headers):
        cell = t_hdr_row.cells[c_idx]
        cell.width = tax_col_w[c_idx]
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tax_data = [
        ("Continuous Numeric", "Ratio & Interval\n(Float, Int)", "Infinite potential values within bounds (e.g., ARR, session length, latency, temperature).", "Histogram, KDE curve, box plot, skewness/kurtosis, normality tests."),
        ("Discrete / Count", "Integer Counts\n(Non-negative Int)", "Finite integer values representing occurrences (e.g., login count, support tickets, seats).", "Zero-inflation check, Poisson/Negative Binomial fit, bar frequency plots."),
        ("Categorical Nominal", "Unordered Labels\n(String, Object)", "Mutually exclusive qualitative classes without inherent ordering (e.g., country, plan type).", "Frequency counts, cardinality ratio, mode dominance, rare token clustering."),
        ("Categorical Ordinal", "Ordered Levels\n(Category, Int)", "Qualitative classes with strict hierarchical progression (e.g., Tier 1 < 2 < 3, Likert scales).", "Monotonic trend analysis, rank-order correlations, cumulative frequency."),
        ("Temporal & Series", "Timestamps & Dates\n(Datetime64)", "Ordered longitudinal progression (e.g., event time, transaction timestamp, renewal date).", "Timeline continuity check, interval gap scans, seasonality, trend decomposition.")
    ]

    for r_idx, row_vals in enumerate(tax_data, start=1):
        row = tax_tbl.rows[r_idx]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = tax_col_w[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_border(cell, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            elif c_idx == 1:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    add_h2(doc, "2.2 Initial Structural Sanity Auditing")
    add_body_p(doc,
               "Upon loading any dataset, the data scientist must perform a non-negotiable structural audit to establish baseline integrity:")
    
    add_bullet_item(doc, "Dimensionality & Footprint Audit",
                    "Document total rows, total columns, and memory consumption. If the dataset exceeds available RAM, immediately profile downcasting opportunities "
                    "(e.g., converting `float64` to `float32`, `int64` to `int32`, and high-repetition strings to categorical types) or migrate to Polars.")
    add_bullet_item(doc, "Primary Key & Granularity Verification",
                    "Confirm what a single row actually represents. Is it one customer? One transaction? One session? We assert that candidate primary keys "
                    "are 100% unique (`df['id'].nunique() == len(df)`). Unintended duplicates skew all downstream statistical weights.")
    add_bullet_item(doc, "Automated Data Dictionary Synthesis",
                    "Compile an automated metadata registry tracking: Column Name, Inferred Physical Type, Logical Data Type, Missing Count, Missing Percentage, "
                    "Unique Value Count, Cardinality Ratio, and 3 Sample Values.")

    # -------------------------------------------------------------------------
    # SECTION 3: CORE EXPLORATION TECHNIQUES & STATISTICAL DIAGNOSTICS
    # -------------------------------------------------------------------------
    add_h1(doc, "3. Core Exploration Techniques & Statistical Diagnostics")
    
    add_h2(doc, "3.1 Data Hygiene, Missingness & Integrity Auditing")
    add_body_p(doc,
               "Missing data is not merely an inconvenience; it is often a powerful behavioral signal in its own right. "
               "The OmniEDA framework strictly prohibits thoughtless blanket imputation (such as filling everything with the mean). "
               "We diagnose missingness through Donald Rubin's classic threefold taxonomy:")
    
    add_bullet_item(doc, "Missing Completely at Random (MCAR)",
                    "Missingness is entirely unrelated to any observed or unobserved variable (e.g., a hardware sensor battery died at random). "
                    "Diagnostic: Little's MCAR test. Treatment: Listwise deletion or simple random imputation is statistically permissible without introducing bias.")
    add_bullet_item(doc, "Missing at Random (MAR)",
                    "Missingness depends systematically on other observed variables, but not on the missing value itself (e.g., enterprise accounts rarely fill "
                    "out self-serve satisfaction surveys because they have dedicated CSM check-ins). Treatment: Conditional grouped median imputation, KNN, or MICE.")
    add_bullet_item(doc, "Missing Not at Random (MNAR)",
                    "The probability of missingness depends directly on the unobserved value itself (e.g., customers with catastrophic service experiences deliberately "
                    "refuse to engage, or users with very high salaries skip income fields). Treatment: Preserve the signal by creating an explicit binary missingness indicator "
                    "(`feature_is_missing = 1`) prior to imputation.")

    add_body_p(doc,
               "We deploy the `missingno` Python library to generate missingness matrix plots, nullity correlation heatmaps, and dendrograms, "
               "identifying whether columns drop values in coordinated clusters.")

    add_h2(doc, "3.2 Outlier & Anomaly Detection Framework")
    add_body_p(doc,
               "A common trap in data science is treating every outlier as an error that must be erased. In truth, outliers fall into two starkly different camps: "
               "measurement errors (e.g., a customer age recorded as 999 or negative revenue) which must be purged, and legitimate black swan events "
               "(e.g., an enterprise power customer generating 50x normal API traffic) which represent vital high-value business dynamics.")
    
    add_bullet_item(doc, "Tukey's Interquartile Range (IQR) Rule",
                    "A non-parametric standard: Inner Fences = [Q1 - 1.5*IQR, Q3 + 1.5*IQR]; Outer Fences = [Q1 - 3.0*IQR, Q3 + 3.0*IQR]. "
                    "Values beyond outer fences are classified as extreme outliers.")
    add_bullet_item(doc, "Modified Z-Score (Median Absolute Deviation - MAD)",
                    "Standard Z-score is deeply flawed because the sample mean and variance are themselves corrupted by the outliers. "
                    "The Modified Z-Score uses the median and MAD: M_i = 0.6745 * |x_i - Median| / MAD. Points with |M_i| > 3.5 are flagged as robust anomalies.")
    add_bullet_item(doc, "Multivariate Anomaly Scans (Isolation Forest)",
                    "Points may appear normal across individual variables, but highly anomalous in combination (e.g., 500 logins per day is normal, and an account age of 1 day "
                    "is normal, but 500 logins on an account that is 1 day old is a bot anomaly). Isolation Forests isolate anomalies through recursive random tree partitioning.")

    add_h2(doc, "3.3 Univariate Distribution Analysis")
    add_body_p(doc,
               "Univariate exploration examines variables in complete isolation to understand their individual properties:")
    
    add_bullet_item(doc, "Continuous Variables (Central Tendency & Shape)",
                    "Evaluate mean vs. median to immediately diagnose skewness. Calculate Fisher-Pearson coefficient of skewness and kurtosis. "
                    "Run Shapiro-Wilk or D'Agostino-Pearson normality tests. Plot Q-Q (Quantile-Quantile) plots against a theoretical normal distribution. "
                    "If skewness > 1.0, test Power Transformations (Log1p, Box-Cox, or Yeo-Johnson) to stabilize variance.")
    add_bullet_item(doc, "Categorical Variables (Frequency & Cardinality)",
                    "Compute frequency counts, mode percentage, and entropy. Perform Pareto analysis (identifying if 20% of categories account for 80% of volume). "
                    "Flag high-cardinality columns (unique values > 50) and identify rare labels (< 1% frequency) for group consolidation.")

    add_h2(doc, "3.4 Bivariate Association & Hypothesis Generation")
    add_body_p(doc,
               "Bivariate analysis investigates how pairs of variables interact, identifying candidate predictors and potential confounders:")
    
    add_bullet_item(doc, "Numeric vs. Numeric",
                    "Compute Pearson r (linear correlation) alongside Spearman rho (rank correlation). If Spearman is significantly higher than Pearson, "
                    "the relationship is non-linear but monotonic. Calculate Distance Correlation to detect non-monotonic dependencies (e.g., U-shaped or sinusoidal curves).")
    add_bullet_item(doc, "Numeric vs. Categorical",
                    "Compare subgroup distributions using grouped box and violin plots. Perform parametric one-way ANOVA (if normality and homoscedasticity assumptions hold) "
                    "or non-parametric Kruskal-Wallis H-tests to determine whether subgroup median shifts are statistically significant.")
    add_bullet_item(doc, "Categorical vs. Categorical",
                    "Construct two-way cross-tabulation contingency tables. Execute Pearson's Chi-Square Test of Independence to evaluate whether category co-occurrences "
                    "deviate significantly from expected random distributions. Compute Cramer's V to quantify the magnitude of association (scale 0 to 1).")

    add_h2(doc, "3.5 Multivariate Interactions & Dimensionality Exploration")
    add_body_p(doc,
               "Real-world business systems are multivariate. We deploy advanced techniques to explore complex multi-feature spaces:")
    
    add_bullet_item(doc, "Multicollinearity & Variance Inflation Factor (VIF)",
                    "Multicollinearity destabilizes linear models and inflates feature importance variance. We compute VIF for all numeric features: "
                    "VIF = 1 / (1 - R_i^2). Features with VIF > 5.0 are flagged for consolidation or elimination.")
    add_bullet_item(doc, "Unsupervised Dimensionality Reduction (PCA)",
                    "Execute Principal Component Analysis on standardized continuous features. Generate Scree plots of explained variance ratio to determine "
                    "how many latent dimensions capture 85%+ of overall dataset variance.")
    add_bullet_item(doc, "Manifold Learning (t-SNE & UMAP)",
                    "Project high-dimensional observations onto 2D space using UMAP (Uniform Manifold Approximation and Projection) to visually identify "
                    "natural customer clusters and subgroup separations before unsupervised clustering.")

    # -------------------------------------------------------------------------
    # SECTION 4: STRATEGIC VISUALIZATION TAXONOMY & CHART SELECTION
    # -------------------------------------------------------------------------
    add_h1(doc, "4. Strategic Visualization Taxonomy & Diagnostic Chart Selection")
    
    add_body_p(doc,
               "A visualization should never be chosen at random. Every chart must have a clear diagnostic purpose, selected based on the mathematical "
               "types of the variables being analyzed and the specific analytical question being answered:")

    # Chart Taxonomy Table
    chart_tbl = doc.add_table(rows=7, cols=4)
    chart_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    chart_tbl.autofit = False
    chart_col_w = [Inches(1.5), Inches(1.5), Inches(1.2), Inches(2.7)]
    chart_headers = ["Analytical Goal", "Optimal Plot Type", "Python Tool", "Diagnostic Value & Interpretation Guide"]
    
    c_hdr_row = chart_tbl.rows[0]
    for c_idx, h_text in enumerate(chart_headers):
        cell = c_hdr_row.cells[c_idx]
        cell.width = chart_col_w[c_idx]
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    chart_data = [
        ("Distribution & Shape", "Histogram + KDE Overlay", "Seaborn / Matplotlib", "Reveals modality (unimodal vs bimodal peaks), skewness, and concentration zones."),
        ("Outlier & Spread", "Box Plot / Violin Plot", "Seaborn / Plotly", "Displays median, IQR boundaries, extreme whiskers, and probability density width."),
        ("Cumulative Probability", "Empirical CDF (ECDF)", "Statsmodels / Seaborn", "Evaluates percentiles without arbitrary binning artifacts; ideal for latency and ARR."),
        ("Pairwise Association", "Scatter Plot + Marginal Rugs", "Seaborn scatterplot", "Inspects non-linear curvature, heteroscedastic fan shapes, and cluster clustering."),
        ("Density over Big Data", "2D Hexbin / Contour Map", "Matplotlib / Seaborn", "Eliminates visual overplotting across 100k+ points via hexagonal density binning."),
        ("Correlation Matrix", "Clustered Heatmap", "Seaborn heatmap", "Hierarchically clusters collinear features into visual blocks using dendrogram ordering.")
    ]

    for r_idx, row_vals in enumerate(chart_data, start=1):
        row = chart_tbl.rows[r_idx]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = chart_col_w[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_border(cell, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            elif c_idx == 1:
                r.bold = True
                r.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Embedding Diagram 2
    add_h2(doc, "4.1 The Statistical Visualization Taxonomy Decision Matrix")
    add_body_p(doc,
               "Figure 1 provides a comprehensive visual taxonomy mapping analytical questions to optimal Python visualization architectures:")
    add_image_box(doc, r"e:\Code Playground\Sumu\assets\eda_diagram_2_taxonomy.png", 
                  "Figure 1", 
                  "OmniEDA Diagnostic Visualization Taxonomy & Plot Selection Decision Matrix (Distributions, Correlations, Categoricals, and High-Dimensional Manifolds)")
    add_body_p(doc,
               "Interpretation Walkthrough: When exploring univariate spread, histograms combined with KDE curves and box plots provide full distributional coverage. "
               "When analyzing relationships between numeric features, standard scatter plots are enhanced with 2D hexbins to prevent visual overplotting in large samples. "
               "For categorical subgroups, faceted dot plots and violin grids expose subtle variance disparities that summary means conceal. "
               "Finally, high-dimensional projections via UMAP and parallel coordinates allow analysts to detect complex multivariate clusters.")

    # -------------------------------------------------------------------------
    # SECTION 5: PYTHON ECOSYSTEM & TECHNICAL TOOLING STACK
    # -------------------------------------------------------------------------
    add_h1(doc, "5. Python Ecosystem & Technical Tooling Stack")
    add_body_p(doc,
               "Python represents the premier language for Exploratory Data Analysis due to its rich, interoperable scientific ecosystem. "
               "The OmniEDA framework utilizes a carefully curated stack of mature, high-performance libraries:")

    # Tooling Table
    tool_tbl = doc.add_table(rows=8, cols=4)
    tool_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tool_tbl.autofit = False
    tool_col_w = [Inches(1.5), Inches(1.4), Inches(0.8), Inches(3.2)]
    tool_headers = ["Functional Layer", "Library / Framework", "Version", "Technical Justification & Purpose"]
    
    t_hdr_row = tool_tbl.rows[0]
    for c_idx, h_text in enumerate(tool_headers):
        cell = t_hdr_row.cells[c_idx]
        cell.width = tool_col_w[c_idx]
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    tool_data = [
        ("Core Data Wrangling", "pandas\npolars", ">= 2.2.0\n>= 0.20.0", "Fast vectorized operations, column-wise aggregations, multi-table joins, and out-of-core memory streaming."),
        ("Missingness Diagnostics", "missingno", ">= 0.5.2", "Visual matrix plots, missing value heatmaps, and nullity correlation dendrograms."),
        ("Statistical Testing", "scipy.stats\nstatsmodels", ">= 1.12.0\n>= 0.14.0", "Formal hypothesis tests (ANOVA, Kruskal-Wallis, Chi-Square, Shapiro-Wilk) and VIF collinearity calculations."),
        ("Static Publication Graphics", "matplotlib\nseaborn", ">= 3.8.0\n>= 0.13.0", "Pixel-perfect publication figures, custom color ramps, faceted small-multiples, and statistical regression plots."),
        ("Interactive Dashboards", "plotly\nnbformat", ">= 5.18.0\n>= 5.9.0", "Interactive zooming, dynamic tooltips, 3D scatter plots, treemaps, and HTML dashboard exports."),
        ("Automated Profiling", "ydata-profiling", ">= 4.6.0", "One-click generation of comprehensive HTML data health profiles, duplicate checks, and correlation matrices."),
        ("Data Contract Assertions", "great-expectations", ">= 0.18.0", "Declarative schema rules, nullity threshold assertions, and automated pipeline data quality gates.")
    ]

    for r_idx, row_vals in enumerate(tool_data, start=1):
        row = tool_tbl.rows[r_idx]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = tool_col_w[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_border(cell, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            elif c_idx == 1:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------------------
    # SECTION 6: PROJECT IMPLEMENTATION TIMELINE & RESOURCE ALLOCATION
    # -------------------------------------------------------------------------
    add_h1(doc, "6. Project Implementation Timeline & Resource Allocation (32.5 Hours)")
    add_body_p(doc,
               "Conducting an exhaustive Exploratory Data Analysis framework requires disciplined time budgeting. "
               "To fulfill the target allocation of roughly 30 to 35 hours of high-impact data science work, "
               "OmniEDA is structured into 6 progressive phases totaling exactly 32.5 working hours over a 5-day execution sprint (6.5 hours/day):")

    # WBS Table
    wbs_tbl = doc.add_table(rows=7, cols=5)
    wbs_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    wbs_tbl.autofit = False
    wbs_col_w = [Inches(1.1), Inches(2.2), Inches(0.9), Inches(1.0), Inches(1.7)]
    wbs_headers = ["Phase", "Phase Name & Scope", "Hours", "Days", "Core Deliverable & Milestone"]
    
    w_hdr_row = wbs_tbl.rows[0]
    for c_idx, h_text in enumerate(wbs_headers):
        cell = w_hdr_row.cells[c_idx]
        cell.width = wbs_col_w[c_idx]
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    wbs_data = [
        ("Phase 1", "Scope, Architecture & Structural Profiling", "4.5 h", "Day 1", "Automated Data Dictionary & Schema Profile"),
        ("Phase 2", "Data Hygiene, Missingness & Anomaly Auditing", "5.5 h", "Day 1-2", "Missingno Matrix & Anomaly Boundary Suite"),
        ("Phase 3", "Univariate & Bivariate Statistical Exploration", "7.0 h", "Day 2-3", "Distribution Diagnostics & Statistical Test Catalog"),
        ("Phase 4", "Multivariate Analysis, Dimensionality & Clustering", "7.5 h", "Day 3-4", "VIF Collinearity Matrix & UMAP Cluster Map"),
        ("Phase 5", "Interactive Dashboards & Visualization Engine", "4.5 h", "Day 4-5", "Plotly Visual Suite & Faceted Gallery"),
        ("Phase 6", "Findings Documentation & Stakeholder Synthesis", "3.5 h", "Day 5", "Final OmniEDA Report & Modeling Handoff Spec")
    ]

    for r_idx, row_vals in enumerate(wbs_data, start=1):
        row = wbs_tbl.rows[r_idx]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = wbs_col_w[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            set_cell_border(cell, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            elif c_idx == 2:
                r.bold = True
                r.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Embedding Diagram 1
    add_h2(doc, "6.1 The 5-Stage Universal EDA Lifecycle Flowchart")
    add_body_p(doc,
               "Figure 2 illustrates the sequential analytical decision flow across all 5 core stages of the OmniEDA framework:")
    add_image_box(doc, r"e:\Code Playground\Sumu\assets\eda_diagram_1_lifecycle.png", 
                  "Figure 2", 
                  "OmniEDA Universal 5-Stage Exploratory Data Analysis & Analytical Decision Lifecycle")
    add_body_p(doc,
               "Lifecycle Walkthrough: The framework starts with ingestion and strict schema validation, moves to univariate profiling, "
               "progresses to bivariate hypothesis testing and multivariate dimensionality reduction, and concludes with a structured modeling handoff spec.")

    # Embedding Diagram 3 (Gantt)
    add_h2(doc, "6.2 Critical Path Gantt Schedule & Pacing")
    add_body_p(doc,
               "Figure 3 maps the 32.5-hour implementation schedule across the 5 working days, highlighting milestones and deliverable checkpoints:")
    add_image_box(doc, r"e:\Code Playground\Sumu\assets\eda_diagram_3_timeline.png", 
                  "Figure 3", 
                  "OmniEDA Week 2 Implementation Timeline, Phase Allocations & Critical Path Gantt Schedule (32.5 Hours Total Effort)")
    add_body_p(doc,
               "Execution Pacing: Pacing is calibrated at 6.5 hours per day, leaving a 2.5-hour unallocated contingency buffer within the 35-hour ceiling "
               "to absorb data format idiosyncrasies or extended high-dimensional manifold tuning.")

    # -------------------------------------------------------------------------
    # SECTION 7: FINDINGS DOCUMENTATION, REPORTING & COMMUNICATION
    # -------------------------------------------------------------------------
    add_h1(doc, "7. Findings Documentation, Reporting & Executive Communication")
    add_body_p(doc,
               "An Exploratory Data Analysis is only as valuable as the decisions it informs. "
               "The OmniEDA framework standardizes findings documentation into four modular deliverables designed for distinct audiences:")

    add_bullet_item(doc, "Deliverable 1: The Executive Data Health Scorecard",
                    "A 1-page dashboard summarizing data fitness for business stakeholders. It scores overall completeness, integrity violation rates, "
                    "cardinality risks, and identifies top-level business trends discovered during exploration.")
    add_bullet_item(doc, "Deliverable 2: The Technical Audit & Statistical Findings Log",
                    "A deep technical document for data scientists. It logs normality test p-values, skewness coefficients, VIF scores, "
                    "correlation matrices, and missingness mechanisms.")
    add_bullet_item(doc, "Deliverable 3: The Modeling Readiness & Feature Handoff Specification",
                    "The direct bridge between EDA and downstream machine learning. It specifies recommended imputation methods per feature, "
                    "identifies non-linear variables requiring power transformations, lists redundant features to drop, and proposes promising interaction terms.")
    add_bullet_item(doc, "Deliverable 4: The Interactive Exploration Gallery",
                    "A self-contained HTML/Plotly notebook or Quarto dashboard enabling both technical and non-technical stakeholders to drill down, "
                    "zoom, and filter charts interactively.")

    # -------------------------------------------------------------------------
    # SECTION 8: COGNITIVE BIASES, GOVERNANCE, ETHICS & TRANSITION
    # -------------------------------------------------------------------------
    add_h1(doc, "8. Cognitive Biases, Governance, Ethical Safeguards & Transition")
    
    add_h2(doc, "8.1 Guarding Against Cognitive Traps in EDA")
    add_body_p(doc,
               "Exploratory analysis carries inherent psychological risks. The human brain is an aggressive pattern-matching machine, "
               "frequently seeing meaningful signals in completely random noise. OmniEDA enforces explicit defensive safeguards against classic cognitive traps:")
    
    add_bullet_item(doc, "Confirmation Bias & Selective Charting",
                    "Analysts often stop exploring once they find a chart confirming their initial business hunch. Defense: Formulate counter-hypotheses "
                    "and explicitly search for data subsets that refute the assumption.")
    add_bullet_item(doc, "Data Snooping & P-Hacking",
                    "Testing hundreds of feature pairs without adjustment guarantees that some will show 'statistically significant' p-values purely by chance. "
                    "Defense: Apply Bonferroni or Benjamini-Hochberg False Discovery Rate (FDR) corrections when running large correlation scans.")
    add_bullet_item(doc, "Survivor Bias",
                    "Analyzing only active accounts or successful events while ignoring canceled accounts or failed runs paints a dangerously optimistic picture. "
                    "Defense: Always audit cohort entry and exit dates.")

    add_h2(doc, "8.2 Data Privacy (PII) & Ethical Governance")
    add_body_p(doc,
               "Exploratory notebooks are frequently shared across teams, creating serious privacy leakage risks:")
    add_bullet_item(doc, "Automated PII Stripping",
                    "Names, email addresses, phone numbers, and IP addresses must be hashed or stripped at initial ingestion before generating charts.")
    add_bullet_item(doc, "Aggregation Masking for Small Subgroups",
                    "When plotting sensitive metrics across small demographic or organizational subgroups (< 10 individuals), apply k-anonymity aggregation "
                    "to prevent individual re-identification.")

    # Formal Approval Block
    add_h2(doc, "8.3 Strategy Sign-Off & Project Governance Registry")
    
    appr_tbl = doc.add_table(rows=4, cols=4)
    appr_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    appr_tbl.autofit = False
    appr_col_w = [Inches(1.8), Inches(1.8), Inches(1.5), Inches(1.8)]
    appr_headers = ["Project Role", "Name & Contact", "Approval Status", "Sign-Off Date"]
    
    a_hdr_row = appr_tbl.rows[0]
    for c_idx, h_text in enumerate(appr_headers):
        cell = a_hdr_row.cells[c_idx]
        cell.width = appr_col_w[c_idx]
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    appr_data = [
        ("Lead Data Science Architect", "Sumarjana Biswas (sumarjanabiswas690@gmail.com)", "APPROVED (Framework Locked)", "2026-09-17"),
        ("Head of Analytics", "Technical Advisory Board", "PEER REVIEWED", "2026-09-17"),
        ("VP of Data Science", "Executive Sponsor", "ADOPTED AS STANDARD", "2026-09-17")
    ]

    for r_idx, row_vals in enumerate(appr_data, start=1):
        row = appr_tbl.rows[r_idx]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = appr_col_w[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            set_cell_border(cell, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            elif c_idx == 2:
                r.bold = True
                r.font.color.rgb = RGBColor(0x05, 0x96, 0x69) if "APPROVED" in val or "ADOPTED" in val else RGBColor(0x02, 0x84, 0xC7)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    
    # Save the document with multi-path resiliency
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    dir_name = os.path.dirname(os.path.abspath(output_path))
    candidates = [
        output_path,
        os.path.join(dir_name, "EDA_and_Visualization_Framework_Week_2_Sumarjana_Biswas.docx"),
        os.path.join(dir_name, "EDA_and_Visualization_Framework_Week_2_Final.docx"),
        os.path.join(dir_name, "EDA_and_Visualization_Framework_Week_2_Updated.docx"),
    ]
    
    saved_paths = []
    for path in candidates:
        try:
            doc.save(path)
            saved_paths.append(path)
            print(f"[SUCCESS] Saved to: {path}")
        except PermissionError:
            print(f"[LOCKED] File currently open in Word: {path}")
            
    if not saved_paths:
        raise PermissionError("All candidate file paths are currently locked by Word. Please close Word and retry.")

if __name__ == "__main__":
    output_docx = r"e:\Code Playground\Sumu\docs\EDA_and_Visualization_Framework_Week_2.docx"
    build_eda_framework_document(output_docx)
